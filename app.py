import streamlit as st
import cv2
import numpy as np
import os
import zipfile
import tempfile
import io
import re
import csv
from pathlib import Path
from math import ceil
from PIL import Image, ImageDraw, ImageFont
import textwrap
import openpyxl


# ─── Helpers ──────────────────────────────────────────────────────────────────

def crop_frame(frame, top, bottom, left, right):
    h, w = frame.shape[:2]
    return frame[int(h*top/100):int(h*(100-bottom)/100), int(w*left/100):int(w*(100-right)/100)]


def resize_frame(frame, max_width):
    h, w = frame.shape[:2]
    if w > max_width:
        scale = max_width / w
        frame = cv2.resize(frame, (max_width, int(h * scale)))
    return frame


def time_to_seconds(t):
    t = t.replace(',', '.').strip()
    parts = t.split(':')
    if len(parts) == 2:
        return int(parts[0]) * 60 + float(parts[1])
    elif len(parts) == 3:
        return int(parts[0]) * 3600 + int(parts[1]) * 60 + float(parts[2])
    return 0


def parse_vtt(content):
    segments = []
    content = content.replace('\r\n', '\n').replace('\r', '\n')
    fillers = {'mhm', 'ja', 'ok', 'okay', 'yes', 'hm', 'hmm', 'uh', 'um', 'ah', 'oh', 'no', 'yeah'}
    blocks = re.split(r'\n[ \t]*\n', content)
    for block in blocks:
        lines = [l.strip() for l in block.strip().split('\n') if l.strip()]
        ts_line_idx = None
        for i, line in enumerate(lines):
            if '-->' in line:
                ts_line_idx = i
                break
        if ts_line_idx is None:
            continue
        ts_match = re.match(
            r'(\d{1,2}:\d{2}:\d{2}[.,]\d{3})\s*-->\s*(\d{1,2}:\d{2}:\d{2}[.,]\d{3})',
            lines[ts_line_idx]
        )
        if not ts_match:
            continue
        start = time_to_seconds(ts_match.group(1))
        end = time_to_seconds(ts_match.group(2))
        text_lines = lines[ts_line_idx + 1:]
        text = ' '.join(re.sub(r'<[^>]+>', '', l) for l in text_lines).strip()
        clean = text.rstrip('.!?,').lower()
        if text and len(text) > 5 and clean not in fillers:
            segments.append((start, end, text))
    return segments


def parse_docx(file_bytes):
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    segments = []
    time_pattern = re.compile(r'\((\d{1,2}:\d{2}(?::\d{2})?)\)')
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    i = 0
    while i < len(paragraphs):
        para = paragraphs[i]
        time_match = time_pattern.search(para)
        if time_match:
            start = time_to_seconds(time_match.group(1))
            texts = []
            i += 1
            while i < len(paragraphs) and not time_pattern.search(paragraphs[i]):
                texts.append(paragraphs[i])
                i += 1
            if texts:
                segments.append((start, start + 15, ' '.join(texts)))
        else:
            i += 1
    return segments


def match_transcript(timestamp_s, segments, window=4):
    texts = [t for s, e, t in segments if s <= timestamp_s + window and e >= timestamp_s - window]
    return ' '.join(texts)


def add_caption_to_image(img_bytes, caption, use_png, jpeg_quality=95, font_scale=4):
    img = Image.open(io.BytesIO(img_bytes)).convert('RGB')
    w, h = img.size
    base_size = max(16, w // 70)
    font_size = base_size * font_scale
    font_paths = [
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
        "/usr/share/fonts/truetype/ubuntu/Ubuntu-R.ttf",
        "/usr/share/fonts/truetype/freefont/FreeSans.ttf",
    ]
    font = None
    for path in font_paths:
        try:
            font = ImageFont.truetype(path, font_size)
            break
        except:
            continue
    if font is None:
        try:
            font = ImageFont.load_default(size=font_size)
        except:
            font = ImageFont.load_default()
            font_size = 10
    char_w = max(1, int(font_size * 0.55))
    chars_per_line = max(10, (w - 24) // char_w)
    wrapped = textwrap.wrap(caption, width=chars_per_line) if caption.strip() else ["(no transcript)"]
    line_h = int(font_size * 1.05)
    box_h = len(wrapped) * line_h + 20
    new_img = Image.new('RGB', (w, h + box_h), (245, 245, 245))
    new_img.paste(img, (0, 0))
    draw = ImageDraw.Draw(new_img)
    y = h + 10
    for line in wrapped:
        draw.text((12, y), line, fill=(30, 30, 30), font=font)
        y += line_h
    out = io.BytesIO()
    if use_png:
        new_img.save(out, format='PNG')
    else:
        new_img.save(out, format='JPEG', quality=jpeg_quality, subsampling=0)
    return out.getvalue()


def get_caption(fname):
    """Read the latest caption — prefer the live widget value over the stored one."""
    return st.session_state.get(f"caption_{fname}", st.session_state.captions.get(fname, ""))


# ─── Session state ─────────────────────────────────────────────────────────────
if 'screenshots' not in st.session_state:
    st.session_state.screenshots = []
if 'captions' not in st.session_state:
    st.session_state.captions = {}
if 'page' not in st.session_state:
    st.session_state.page = 0
if 'deselected' not in st.session_state:
    st.session_state.deselected = set()


# ─── UI ───────────────────────────────────────────────────────────────────────
st.set_page_config(page_title="Video → Screenshots", layout="centered")
st.title("Video → Screenshots")

st.header("Step 1: Extract Screenshots")

col1, col2 = st.columns(2)
with col1:
    interval = st.selectbox(
        "Interval between screenshots",
        options=[1, 2, 3, 4, 5],
        index=2,
        format_func=lambda x: f"every {x} second{'s' if x > 1 else ''}",
    )
with col2:
    split_parts = st.selectbox(
        "Split video into",
        options=[1, 2, 3, 4],
        index=0,
        format_func=lambda x: "no splitting" if x == 1 else f"{x} parts",
    )

st.caption("Tip: For large videos (over 500 MB), split into 2–3 parts to keep each download manageable.")

with st.expander("Quality settings"):
    quality_preset = st.radio(
        "Screenshot quality",
        options=["Standard (faster, smaller files)", "High (recommended)", "Maximum (slowest, largest files)"],
        index=2,
    )
    quality_map = {
        "Standard (faster, smaller files)": (800, 85),
        "High (recommended)": (1280, 95),
        "Maximum (slowest, largest files)": (1920, 100),
    }
    max_width, jpeg_quality = quality_map[quality_preset]
    img_format = st.radio(
        "Image format",
        options=["JPEG (smaller files, faster)", "PNG (lossless, sharper text)"],
        index=1,
    )
    use_png = "PNG" in img_format

with st.expander("Crop settings (recommended: remove browser bar, macOS bar & webcam panel)"):
    st.markdown("Cut away the macOS menu bar, browser bar, and webcam panel on the right so only the shared screen content is kept.")
    use_crop = st.checkbox("Enable crop")
    if use_crop:
        crop_col1, crop_col2 = st.columns(2)
        with crop_col1:
            crop_top = st.slider("Cut from top (%)", 0, 40, 19)
            crop_left = st.slider("Cut from left (%)", 0, 40, 0)
        with crop_col2:
            crop_bottom = st.slider("Cut from bottom (%)", 0, 40, 5)
            crop_right = st.slider("Cut from right (%)", 0, 40, 13)
    else:
        crop_top = crop_bottom = crop_left = crop_right = 0

with st.expander("Combined Screenshot Layout (recommended for Mural uploads)"):
    st.markdown("Instead of hundreds of individual images, screenshots are combined into wide rows (e.g. 5 per row). This keeps them in the correct left-to-right order when uploading to tools like **Mural**.")
    use_combined_layout = st.checkbox("Create Combined Screenshot Layout instead of individual screenshots")
    if use_combined_layout:
        per_row = st.select_slider("Screenshots per row", options=[3, 4, 5, 6], value=5)
        add_spacing = st.checkbox("Add spacing between screenshots", value=True)
        if add_spacing:
            spacing = st.slider("Spacing between screenshots (pixels)", min_value=5, max_value=60, value=45, step=5)
        else:
            spacing = 0
    else:
        per_row = 5
        spacing = 0

uploaded_file = st.file_uploader("Upload video", type=["mp4", "mov", "avi", "mkv"])

if uploaded_file:
    file_size_mb = len(uploaded_file.getvalue()) / (1024 * 1024)
    st.info(f"**{uploaded_file.name}** — {file_size_mb:.1f} MB")

    if st.button("Create screenshots", type="primary"):
        suffix = Path(uploaded_file.name).suffix or ".mp4"
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            tmp.write(uploaded_file.getvalue())
            tmp_path = tmp.name

        try:
            cap = cv2.VideoCapture(tmp_path)
            if not cap.isOpened():
                st.error("Could not open video. Please try a different format.")
                st.stop()

            fps = cap.get(cv2.CAP_PROP_FPS) or 25
            total_frames = int(cap.get(cv2.CAP_PROP_FRAME_COUNT))
            duration_s = total_frames / fps

            st.write(f"Duration: **{duration_s/60:.1f} min** | FPS: **{fps:.1f}** | Expected: **~{int(duration_s/interval)} screenshots**")

            interval_frames = max(1, int(interval * fps))
            frames_per_part = ceil(total_frames / split_parts)
            progress_bar = st.progress(0.0)
            status = st.empty()

            parts = {p: [] for p in range(1, split_parts + 1)}
            all_screenshots = []
            screenshot_count = 0
            skipped_count = 0
            frame_number = 0
            last_saved_gray = None

            while True:
                ret, frame = cap.read()
                if not ret:
                    break

                if frame_number % interval_frames == 0:
                    processed = crop_frame(frame, crop_top, crop_bottom, crop_left, crop_right) if use_crop else frame
                    processed = resize_frame(processed, max_width=max_width)
                    gray = cv2.cvtColor(processed, cv2.COLOR_BGR2GRAY)
                    is_duplicate = last_saved_gray is not None and np.mean(cv2.absdiff(last_saved_gray, gray)) < 1.5

                    if not is_duplicate:
                        part_num = min(frame_number // frames_per_part + 1, split_parts)
                        ts = frame_number / fps
                        ts_str = f"{int(ts//3600):02d}h{int((ts%3600)//60):02d}m{int(ts%60):02d}s"
                        ext = "png" if use_png else "jpg"
                        fname = f"screenshot_{screenshot_count:04d}_{ts_str}.{ext}"
                        if use_png:
                            _, img_bytes = cv2.imencode(".png", processed)
                        else:
                            _, img_bytes = cv2.imencode(".jpg", processed, [cv2.IMWRITE_JPEG_QUALITY, jpeg_quality])
                        img_bytes = img_bytes.tobytes()
                        parts[part_num].append((fname, img_bytes, ts))
                        all_screenshots.append((fname, img_bytes, ts))
                        screenshot_count += 1
                        last_saved_gray = gray
                    else:
                        skipped_count += 1

                frame_number += 1
                if frame_number % 300 == 0:
                    progress_bar.progress(min(frame_number / total_frames, 1.0))
                    status.text(f"{frame_number/total_frames*100:.0f}% processed — {screenshot_count} screenshots, {skipped_count} duplicates skipped")

            cap.release()
            progress_bar.progress(1.0)
            status.text(f"Done! {screenshot_count} screenshots created, {skipped_count} duplicates skipped.")
            st.success(f"{screenshot_count} screenshots are ready!")

            st.session_state.screenshots = all_screenshots
            st.session_state.captions = {fname: "" for fname, _, _ in all_screenshots}
            st.session_state.deselected = set()
            st.session_state.page = 0

            for part_num in range(1, split_parts + 1):
                part_screenshots = parts[part_num]
                if not part_screenshots:
                    continue

                if use_combined_layout:
                    sheets_buffer = io.BytesIO()
                    sheet_count = 0
                    with zipfile.ZipFile(sheets_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                        for i in range(0, len(part_screenshots), per_row):
                            row_data = part_screenshots[i:i+per_row]
                            row_frames = []
                            for _, img_bytes, _ in row_data:
                                arr = np.frombuffer(img_bytes, np.uint8)
                                frm = cv2.imdecode(arr, cv2.IMREAD_COLOR)
                                if frm is not None:
                                    row_frames.append(frm)
                            if not row_frames:
                                continue
                            # Normalize all frames to the same height
                            target_h = row_frames[0].shape[0]
                            normalized = []
                            for frm in row_frames:
                                if frm.shape[0] != target_h:
                                    scale = target_h / frm.shape[0]
                                    frm = cv2.resize(frm, (int(frm.shape[1] * scale), target_h))
                                normalized.append(frm)
                            # Pad last row with white frames if needed
                            while len(normalized) < per_row:
                                blank = np.ones_like(normalized[0]) * 240
                                normalized.append(blank)
                            if spacing > 0:
                                sep = np.ones((target_h, spacing, 3), dtype=np.uint8) * 240
                                interleaved = []
                                for idx, frm in enumerate(normalized):
                                    interleaved.append(frm)
                                    if idx < len(normalized) - 1:
                                        interleaved.append(sep)
                                sheet = np.hstack(interleaved)
                            else:
                                sheet = np.hstack(normalized)
                            if use_png:
                                _, sheet_bytes = cv2.imencode(".png", sheet)
                                zf.writestr(f"combined_layout_{sheet_count+1:03d}.png", sheet_bytes.tobytes())
                            else:
                                _, sheet_bytes = cv2.imencode(".jpg", sheet, [cv2.IMWRITE_JPEG_QUALITY, jpeg_quality])
                                zf.writestr(f"combined_layout_{sheet_count+1:03d}.jpg", sheet_bytes.tobytes())
                            sheet_count += 1
                    sheets_buffer.seek(0)
                    part_label = f" (part {part_num} of {split_parts})" if split_parts > 1 else ""
                    st.download_button(
                        f"Download {sheet_count} combined layouts{part_label} as ZIP",
                        sheets_buffer,
                        file_name=f"combined_layout_part{part_num}.zip" if split_parts > 1 else "combined_layout.zip",
                        mime="application/zip",
                        key=f"sheet_{part_num}",
                    )
                else:
                    zip_buffer = io.BytesIO()
                    with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                        for fname, img_bytes, _ in part_screenshots:
                            if use_png:
                                arr = np.frombuffer(img_bytes, np.uint8)
                                frm = cv2.imdecode(arr, cv2.IMREAD_COLOR)
                                _, png_bytes = cv2.imencode(".png", frm)
                                zf.writestr(fname, png_bytes.tobytes())
                            else:
                                zf.writestr(fname, img_bytes)
                    zip_buffer.seek(0)
                    label = f"Download part {part_num} of {split_parts} ({len(part_screenshots)} screenshots)" if split_parts > 1 else f"Download {screenshot_count} screenshots as ZIP"
                    st.download_button(
                        label, zip_buffer,
                        file_name=f"screenshots_part{part_num}.zip" if split_parts > 1 else "screenshots.zip",
                        mime="application/zip",
                        key=f"part_{part_num}",
                    )

        finally:
            os.unlink(tmp_path)


# ─── Step 2: Review & Deselect Screenshots ────────────────────────────────────
if st.session_state.screenshots:
    st.divider()
    st.header("Step 2: Review Screenshots")
    st.markdown("Uncheck screenshots you don't need — they will be excluded from all downloads.")

    screenshots = st.session_state.screenshots
    total = len(screenshots)
    deselected = st.session_state.deselected

    sel_col1, sel_col2, sel_col3 = st.columns([1, 1, 2])
    with sel_col1:
        if st.button("Select all"):
            st.session_state.deselected = set()
            st.rerun()
    with sel_col2:
        if st.button("Deselect all"):
            st.session_state.deselected = {fname for fname, _, _ in screenshots}
            st.rerun()
    with sel_col3:
        kept = total - len(deselected)
        st.markdown(f"<div style='padding-top:8px'>{kept} of {total} selected</div>", unsafe_allow_html=True)

    total_pages = ceil(total / 10)
    page = st.session_state.page

    col_prev, col_info, col_next = st.columns([1, 2, 1])
    with col_prev:
        if st.button("← Previous", key="prev_step2", disabled=page == 0):
            st.session_state.page -= 1
            st.rerun()
    with col_info:
        st.markdown(f"<div style='text-align:center'>Page {page+1} of {total_pages}</div>", unsafe_allow_html=True)
    with col_next:
        if st.button("Next →", key="next_step2", disabled=page >= total_pages - 1):
            st.session_state.page += 1
            st.rerun()

    start_idx = page * 10
    end_idx = min(start_idx + 10, total)

    for fname, img_bytes, ts in screenshots[start_idx:end_idx]:
        ts_str = f"{int(ts//3600):02d}:{int((ts%3600)//60):02d}:{int(ts%60):02d}"
        col_cb, col_img = st.columns([1, 6])
        with col_cb:
            checked = st.checkbox("Keep", value=(fname not in deselected), key=f"sel_{fname}", label_visibility="collapsed")
            if not checked:
                st.session_state.deselected.add(fname)
            else:
                st.session_state.deselected.discard(fname)
        with col_img:
            arr = np.frombuffer(img_bytes, np.uint8)
            thumb = cv2.imdecode(arr, cv2.IMREAD_COLOR)
            thumb = resize_frame(thumb, max_width=800)
            _, thumb_bytes = cv2.imencode(".jpg", thumb, [cv2.IMWRITE_JPEG_QUALITY, 88])
            st.image(thumb_bytes.tobytes(), caption=ts_str, use_container_width=True)
        st.divider()

    selected_screenshots = [(f, b, t) for f, b, t in screenshots if f not in st.session_state.deselected]
    if selected_screenshots:
        if st.button(f"Generate ZIP ({len(selected_screenshots)} selected screenshots)", key="gen_selected"):
            with st.spinner("Creating ZIP..."):
                zip_buffer = io.BytesIO()
                with zipfile.ZipFile(zip_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                    for fname, img_bytes, _ in selected_screenshots:
                        if use_png:
                            arr = np.frombuffer(img_bytes, np.uint8)
                            frm = cv2.imdecode(arr, cv2.IMREAD_COLOR)
                            _, png_bytes = cv2.imencode(".png", frm)
                            zf.writestr(fname, png_bytes.tobytes())
                        else:
                            zf.writestr(fname, img_bytes)
                zip_buffer.seek(0)
            st.download_button(
                f"Download {len(selected_screenshots)} selected screenshots as ZIP",
                zip_buffer,
                file_name="screenshots_selected.zip",
                mime="application/zip",
                key="dl_selected",
            )

        st.markdown("**Download as Combined Screenshot Layout**")
        csl_col1, csl_col2 = st.columns(2)
        with csl_col1:
            csl_per_row = st.select_slider("Screenshots per row", options=[3, 4, 5, 6], value=5, key="csl_per_row")
        with csl_col2:
            csl_spacing = st.slider("Spacing (pixels)", min_value=0, max_value=60, value=45, step=5, key="csl_spacing")

        if st.button("Generate Combined Screenshot Layout", key="gen_csl"):
            with st.spinner("Combining screenshots..."):
                csl_buffer = io.BytesIO()
                sheet_count = 0
                with zipfile.ZipFile(csl_buffer, "w", zipfile.ZIP_DEFLATED) as zf:
                    for i in range(0, len(selected_screenshots), csl_per_row):
                        row_data = selected_screenshots[i:i+csl_per_row]
                        row_frames = []
                        for _, img_bytes, _ in row_data:
                            arr = np.frombuffer(img_bytes, np.uint8)
                            frm = cv2.imdecode(arr, cv2.IMREAD_COLOR)
                            if frm is not None:
                                row_frames.append(frm)
                        if not row_frames:
                            continue
                        target_h = row_frames[0].shape[0]
                        normalized = []
                        for frm in row_frames:
                            if frm.shape[0] != target_h:
                                scale = target_h / frm.shape[0]
                                frm = cv2.resize(frm, (int(frm.shape[1] * scale), target_h))
                            normalized.append(frm)
                        while len(normalized) < csl_per_row:
                            blank = np.ones_like(normalized[0]) * 240
                            normalized.append(blank)
                        if csl_spacing > 0:
                            sep = np.ones((target_h, csl_spacing, 3), dtype=np.uint8) * 240
                            interleaved = []
                            for idx, frm in enumerate(normalized):
                                interleaved.append(frm)
                                if idx < len(normalized) - 1:
                                    interleaved.append(sep)
                            sheet = np.hstack(interleaved)
                        else:
                            sheet = np.hstack(normalized)
                        if use_png:
                            _, sheet_bytes = cv2.imencode(".png", sheet)
                            zf.writestr(f"combined_layout_{sheet_count+1:03d}.png", sheet_bytes.tobytes())
                        else:
                            _, sheet_bytes = cv2.imencode(".jpg", sheet, [cv2.IMWRITE_JPEG_QUALITY, jpeg_quality])
                            zf.writestr(f"combined_layout_{sheet_count+1:03d}.jpg", sheet_bytes.tobytes())
                        sheet_count += 1
                csl_buffer.seek(0)
            st.download_button(
                f"Download {sheet_count} combined layouts as ZIP",
                csl_buffer,
                file_name="combined_layout_selected.zip",
                mime="application/zip",
                key="dl_csl",
            )


# ─── Step 3: Transcript Matching ──────────────────────────────────────────────
if st.session_state.screenshots:
    st.divider()
    st.header("Step 3: Transcript Matching")
    st.markdown("Upload your meeting transcript to automatically match each screenshot with the spoken context. You can edit the text before downloading.")

    transcript_file = st.file_uploader(
        "Upload transcript (.vtt or .docx)",
        type=["vtt", "docx"],
        key="transcript_upload",
    )

    if transcript_file:
        file_bytes = transcript_file.read()
        if transcript_file.name.lower().endswith(".vtt"):
            segments = parse_vtt(file_bytes.decode("utf-8", errors="ignore"))
        else:
            try:
                segments = parse_docx(file_bytes)
            except Exception as e:
                st.error(f"Could not read DOCX file: {e}")
                segments = []

        if not segments:
            st.warning("Could not parse transcript. Please check the file format.")
        else:
            st.success(f"Transcript loaded: {len(segments)} segments found.")

            screenshots = st.session_state.screenshots
            active = [(f, b, t) for f, b, t in screenshots if f not in st.session_state.deselected]

            if all(v == "" for v in st.session_state.captions.values()):
                for fname, _, ts in active:
                    st.session_state.captions[fname] = match_transcript(ts, segments)

            total_pages = ceil(len(active) / 10)
            page = st.session_state.page

            col_prev, col_info, col_next = st.columns([1, 2, 1])
            with col_prev:
                if st.button("← Previous", key="prev_step3", disabled=page == 0):
                    st.session_state.page -= 1
                    st.rerun()
            with col_info:
                st.markdown(f"<div style='text-align:center'>Page {page+1} of {total_pages} &nbsp;|&nbsp; {len(active)} screenshots</div>", unsafe_allow_html=True)
            with col_next:
                if st.button("Next →", key="next_step3", disabled=page >= total_pages - 1):
                    st.session_state.page += 1
                    st.rerun()

            start = page * 10
            end = min(start + 10, len(active))

            for fname, img_bytes, ts in active[start:end]:
                ts_str = f"{int(ts//3600):02d}:{int((ts%3600)//60):02d}:{int(ts%60):02d}"
                col_img, col_text = st.columns([1, 2])
                with col_img:
                    arr = np.frombuffer(img_bytes, np.uint8)
                    thumb = cv2.imdecode(arr, cv2.IMREAD_COLOR)
                    thumb = resize_frame(thumb, max_width=800)
                    _, thumb_bytes = cv2.imencode(".jpg", thumb, [cv2.IMWRITE_JPEG_QUALITY, 88])
                    st.image(thumb_bytes.tobytes(), caption=ts_str, use_container_width=True)
                with col_text:
                    new_text = st.text_area(
                        f"caption_{fname}",
                        value=st.session_state.captions.get(fname, ""),
                        height=120,
                        key=f"caption_{fname}",
                        label_visibility="collapsed",
                    )
                    st.session_state.captions[fname] = new_text
                st.divider()

            st.subheader("Download")

            font_scale = st.slider(
                "Caption text size",
                min_value=1, max_value=8, value=4,
                help="Controls how large the text appears in the caption bar below each image."
            )

            dl_col1, dl_col2, dl_col3 = st.columns(3)

            with dl_col1:
                if st.button("Generate images with caption bar", key="gen_annotated"):
                    with st.spinner("Adding captions to images..."):
                        zip_buf = io.BytesIO()
                        with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
                            for fname, img_bytes, _ in active:
                                caption = get_caption(fname)
                                annotated = add_caption_to_image(img_bytes, caption, use_png, jpeg_quality=jpeg_quality, font_scale=font_scale)
                                ext = "png" if use_png else "jpg"
                                zf.writestr(fname.rsplit(".", 1)[0] + f"_annotated.{ext}", annotated)
                        zip_buf.seek(0)
                    st.download_button(
                        "Download annotated images ZIP",
                        zip_buf,
                        file_name="screenshots_annotated.zip",
                        mime="application/zip",
                        key="dl_annotated",
                    )

            with dl_col2:
                csv_buf = io.StringIO()
                writer = csv.writer(csv_buf)
                writer.writerow(["Screenshot", "Description"])
                for fname, _, _ in active:
                    caption = get_caption(fname).replace('\n', ' ').replace('\r', ' ')
                    writer.writerow([fname, caption])
                st.download_button(
                    "Download CSV",
                    csv_buf.getvalue(),
                    file_name="transcript_mapping.csv",
                    mime="text/csv",
                    key="dl_csv",
                )

            with dl_col3:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Screenshots"
                ws.append(["Screenshot", "Description"])
                ws.column_dimensions["A"].width = 40
                ws.column_dimensions["B"].width = 80
                for fname, _, _ in active:
                    caption = get_caption(fname).replace('\n', ' ').replace('\r', ' ')
                    ws.append([fname, caption])
                excel_buf = io.BytesIO()
                wb.save(excel_buf)
                excel_buf.seek(0)
                st.download_button(
                    "Download Excel",
                    excel_buf,
                    file_name="transcript_mapping.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    key="dl_excel",
                )
